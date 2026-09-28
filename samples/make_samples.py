"""Generates messy sample spreadsheets for trying DataBridge."""
from datetime import date
from pathlib import Path
from openpyxl import Workbook

HERE = Path(__file__).parent

def vendor_sales():
    wb = Workbook(); ws = wb.active; ws.title = "Sales"
    ws["A1"] = "ACME Corp Monthly Sales Report"; ws.merge_cells("A1:G1")
    ws["A2"] = "Generated 2026-09-01"
    ws.append([])
    ws.append(["Cust ID", "First Name", "Last Name", "Region", "Order Date", "Amount", "Status"])
    rows = [
        ["00123", "Jane", "Doe", "East", date(2026, 8, 3), 1200.5, "Shipped"],
        ["00124", "Ravi", "Kumar", None, date(2026, 8, 5), 310, "Pending"],
        ["00125", " maria ", "Lopez", "West", date(2026, 8, 9), "1,045.00", "Shipped"],
        ["00126", "Tom", "Nguyen", "North", date(2026, 8, 12), 87.25, "Cancelled"],
        ["00127", "Aisha", "Bello", "South", date(2026, 8, 15), 2250, "shipped"],
        [None, None, None, None, None, None, None],
        ["00128", "Li", "Wei", "East", date(2026, 8, 21), "n/a", "Pending"],
    ]
    for r in rows: ws.append(r)
    ws.merge_cells("D5:D6")  # East spans two customers
    ws.append(["Total", None, None, None, None, 4892.75, None])
    for r in range(5, 13): ws.cell(r, 5).number_format = "yyyy-mm-dd"
    reg = wb.create_sheet("Regions"); reg.append(["Code", "Region Name"])
    for c, n in [("East", "Eastern US"), ("West", "Western US"), ("North", "Northern US"), ("South", "Southern US")]: reg.append([c, n])
    wb.save(HERE / "vendor_sales.xlsx")

def target_template():
    wb = Workbook(); ws = wb.active; ws.title = "Customers"
    ws.append(["customer_id", "full_name", "region", "order_date", "amount_usd", "status"])
    wb.save(HERE / "target_customer_orders.xlsx")

if __name__ == "__main__":
    vendor_sales(); target_template(); print("samples written")
