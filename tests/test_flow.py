"""End to end: upload -> profile -> target -> map -> publish -> endpoint -> REST call -> re-ingest."""
import io
import sqlite3

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from databridge.api.runtime import router
from databridge.engine.mapper import MappingSpec
from databridge.ingest.sheet_profile import suggest_profile
from databridge.services import connections, endpoints, mappings, sources, targets

app = FastAPI()
app.include_router(router, prefix="/api/v1")
client = TestClient(app)


def test_spreadsheet_to_endpoint(sales_bytes, template_bytes, tmp_path):
    prof = suggest_profile(sales_bytes, "vendor_sales.xlsx")
    src = sources.create_upload_source("Vendor sales", "vendor_sales.xlsx", sales_bytes, prof)
    assert src.latest_snapshot_id and len(src.fields) == 7

    tgt = targets.target_from_template("Customer orders", "target.xlsx", template_bytes)
    targets.save_target(tgt.name, [
        {"name": "customer_id", "type": "string", "required": True},
        {"name": "full_name", "type": "string"},
        {"name": "region", "type": "string"},
        {"name": "order_date", "type": "date"},
        {"name": "amount_usd", "type": "decimal(12,2)", "required": True},
        {"name": "status", "type": "string"},
    ], target_id=tgt.id)

    m = mappings.create_mapping("Sales to orders", src.id, tgt.id)
    auto_targets = {r["target"] for r in m.rules}
    assert {"customer_id", "region", "order_date", "status", "amount_usd"} <= auto_targets
    rules = m.rules + [{"target": "full_name", "formula": 'PROPER([First Name]) & " " & [Last Name]'}]
    mappings.save_spec(m.id, MappingSpec(rules=rules))
    prev = mappings.preview(m.id)
    assert prev.stats == {"rows_in": 6, "rows_out": 5, "rows_rejected": 1}

    ds = mappings.publish(m.id)
    assert ds.version == 1 and ds.row_count == 5 and ds.rejected_count == 1

    endpoints.save_endpoint({"slug": "customer-orders", "name": "Customer orders", "mapping_id": m.id,
                             "params": [{"name": "region", "column": "region", "op": "eq"},
                                        {"name": "q", "column": "full_name", "op": "contains"}]})
    key, raw = endpoints.create_api_key("tests")

    assert client.get("/api/v1/data/customer-orders").status_code == 401
    r = client.get("/api/v1/data/customer-orders", params={"region": "East"}, headers={"X-API-Key": raw})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["total"] == 2 and {d["customer_id"] for d in body["data"]} == {"00123", "00124"}
    r = client.get("/api/v1/data/customer-orders", params={"q": "lopez"}, headers={"X-API-Key": raw})
    assert r.json()["data"][0]["full_name"] == "Maria Lopez"
    r = client.get("/api/v1/data/customer-orders", params={"format": "csv"}, headers={"X-API-Key": raw})
    assert r.text.splitlines()[0].startswith("customer_id,full_name")
    r = client.get("/api/v1/data/customer-orders", params={"format": "xlsx"}, headers={"X-API-Key": raw})
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.active.max_row == 6
    assert client.get("/api/v1/data/customer-orders/schema", headers={"X-API-Key": raw}).json()["fields"][0]["name"] == "customer_id"

    # Same file again is skipped; a changed file publishes v2 automatically.
    r = client.post(f"/api/v1/ingest/{src.id}", files={"file": ("vendor_sales.xlsx", sales_bytes)},
                    headers={"X-API-Key": raw})
    assert r.json()["skipped"] is True
    wb = load_workbook(io.BytesIO(sales_bytes))
    ws = wb["Sales"]
    ws.cell(11, 6).value = 99  # fill the n/a amount
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"/api/v1/ingest/{src.id}", files={"file": ("vendor_sales_v2.xlsx", buf.getvalue())},
                    headers={"X-API-Key": raw})
    out = r.json()
    assert out["published"][0]["version"] == 2 and out["published"][0]["rows"] == 6
    run_id = out["published"][0]["run_id"]
    assert client.get(f"/api/v1/runs/{run_id}", headers={"X-API-Key": raw}).json()["kind"] == "publish"

    # Scripts look the id up by name; the call to use depends on the source kind.
    found = client.get("/api/v1/sources", params={"name": src.name}, headers={"X-API-Key": raw}).json()
    assert found == [{"id": src.id, "name": src.name, "kind": "upload", "latest_snapshot_id": found[0]["latest_snapshot_id"],
                      "load_with": f"/api/v1/ingest/{src.id}"}]
    assert client.get("/api/v1/sources").status_code == 401


def test_database_source(tmp_path):
    db = tmp_path / "erp.db"
    con = sqlite3.connect(db)
    con.executescript("""
        create table orders(order_id integer primary key, amount numeric(12,2), status text);
        insert into orders values (1, 10.5, 'Shipped'), (2, 20, 'Pending');
    """)
    con.commit()
    con.close()
    conn = connections.save_connection("ERP", "database", {"dialect": "sqlite", "database": str(db)})
    assert connections.test_connection(conn.id).ok
    tables = connections.browse(conn.id, {"schema": "main"})
    assert [t.name for t in tables] == ["orders"]
    src = sources.create_connector_source("ERP orders", conn.id, {"schema": "main", "table": "orders"})
    assert sources.load_snapshot(src.id).height == 2
    q = sources.create_connector_source("ERP totals", conn.id, {"sql": "select status, sum(amount) total from orders group by status"})
    assert sources.load_snapshot(q.id).height == 2
